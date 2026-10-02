from docx import Document
from docx.shared import Inches, Pt
from docx.enum.text import WD_ALIGN_PARAGRAPH
from io import BytesIO

def format_docx(document_text, title="LegalEase Document"):
    if not document_text or not document_text.strip():
        raise ValueError("Document content is empty")

    doc = Document()
    section = doc.sections[0]
    section.top_margin = Inches(0.8)
    section.bottom_margin = Inches(0.8)
    section.left_margin = Inches(1)
    section.right_margin = Inches(1)

    style = doc.styles["Normal"]
    style.font.name = "Times New Roman"
    style.font.size = Pt(12)

    heading = doc.add_heading(title, level=0)
    heading.alignment = WD_ALIGN_PARAGRAPH.CENTER

    for line in document_text.splitlines():
        line = line.strip()
        if not line:
            continue
        if line.isupper() and len(line) < 90:
            doc.add_heading(line, level=1)
        else:
            doc.add_paragraph(line)

    footer = section.footer.paragraphs[0]
    footer.alignment = WD_ALIGN_PARAGRAPH.CENTER
    footer.add_run("LegalEase | AI-generated draft – Professional review required")

    output = BytesIO()
    doc.save(output)
    return output.getvalue()